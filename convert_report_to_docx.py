import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def sanitize_xml_string(s):
    """Strip XML 1.0 incompatible control characters."""
    if not isinstance(s, str):
        return s
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', s)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twips (1/20 of a pt). 100 twips = 5pt."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F2F2F2"):
    """Set cell background shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def add_page_number_to_run(run):
    """Inserts a dynamic Word PAGE field into a run."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def parse_inline_tokens(text):
    """
    Parses a string for markdown inline syntax:
    - ***bold_italic***
    - **bold** or <b>bold</b>
    - *italic* or <i>italic</i>
    - `code`
    - $math$
    Returns a list of (style_tag, text_fragment) tuples.
    """
    # Replace html <b> and <i> tags with markdown equivalents
    text = re.sub(r'<b>(.*?)</b>', r'**\1**', text, flags=re.IGNORECASE)
    text = re.sub(r'<i>(.*?)</i>', r'*\1*', text, flags=re.IGNORECASE)
    # Strip <br/> or <br>
    text = re.sub(r'<br\s*/?>', ' ', text, flags=re.IGNORECASE)
    # Replace common HTML entities
    text = text.replace('&nbsp;', ' ')
    text = text.replace('&amp;', '&')
    text = text.replace('&lt;', '<')
    text = text.replace('&gt;', '>')

    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$)')
    tokens = []
    last_end = 0
    for m in pattern.finditer(text):
        start, end = m.span()
        if start > last_end:
            tokens.append(('normal', text[last_end:start]))
        matched = m.group(0)
        if matched.startswith('***') and matched.endswith('***') and len(matched) >= 6:
            tokens.append(('bold_italic', matched[3:-3]))
        elif matched.startswith('**') and matched.endswith('**') and len(matched) >= 4:
            tokens.append(('bold', matched[2:-2]))
        elif matched.startswith('*') and matched.endswith('*') and len(matched) >= 2:
            tokens.append(('italic', matched[1:-1]))
        elif matched.startswith('`') and matched.endswith('`') and len(matched) >= 2:
            tokens.append(('code', matched[1:-1]))
        elif matched.startswith('$') and matched.endswith('$') and len(matched) >= 2:
            tokens.append(('math', matched[1:-1]))
        last_end = end
    if last_end < len(text):
        tokens.append(('normal', text[last_end:]))
    return tokens

def add_runs_to_paragraph(paragraph, tokens, default_font="Times New Roman", default_size=12, default_color=None):
    """Adds styled runs to a paragraph based on inline tokens."""
    for style_tag, token_text in tokens:
        token_text = sanitize_xml_string(token_text)
        if not token_text:
            continue
        run = paragraph.add_run(token_text)
        run.font.name = default_font
        run.font.size = Pt(default_size)
        if default_color:
            run.font.color.rgb = default_color
            
        if style_tag == 'bold':
            run.bold = True
        elif style_tag == 'italic':
            run.italic = True
        elif style_tag == 'bold_italic':
            run.bold = True
            run.italic = True
        elif style_tag == 'code':
            run.font.name = 'Consolas'
            run.font.size = Pt(default_size - 1.5 if default_size > 10 else default_size)
            run.font.color.rgb = RGBColor(0x99, 0x00, 0x00) # dark reddish code
        elif style_tag == 'math':
            run.italic = True
            run.font.name = 'Cambria Math'

def add_signature_table(doc, left_lines, right_lines):
    """Creates a 2-column borderless signature table."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # remove borders
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl._tbl.tblPr.append(tblBorders)
    
    # Left cell
    c0 = tbl.cell(0, 0)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_after = Pt(2)
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for idx, l in enumerate(left_lines):
        if idx > 0:
            p0 = c0.add_paragraph()
            p0.paragraph_format.line_spacing = 1.15
            p0.paragraph_format.space_after = Pt(2)
            p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tokens = parse_inline_tokens(l)
        add_runs_to_paragraph(p0, tokens, default_size=11)
        
    # Right cell
    c1 = tbl.cell(0, 1)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(2)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for idx, l in enumerate(right_lines):
        if idx > 0:
            p1 = c1.add_paragraph()
            p1.paragraph_format.line_spacing = 1.15
            p1.paragraph_format.space_after = Pt(2)
            p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        tokens = parse_inline_tokens(l)
        add_runs_to_paragraph(p1, tokens, default_size=11)

    # Post table spacing
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0

def build_docx_report(md_path, docx_path):
    print(f"Reading markdown from {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = sanitize_xml_string(f.read())

    doc = docx.Document()
    
    # Page Setup (A4 and VIT Margins)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.81)    # 3.81 cm
    section.right_margin = Cm(2.54)   # 2.54 cm
    section.top_margin = Cm(2.54)     # 2.54 cm
    section.bottom_margin = Cm(2.54)  # 2.54 cm

    # Add Page numbering to footer (centered)
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run()
    footer_run.font.name = 'Times New Roman'
    footer_run.font.size = Pt(10)
    add_page_number_to_run(footer_run)

    # Configure Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    # Strip HTML comments like <!-- ... -->
    md_text = re.sub(r'<!--[\s\S]*?-->', '', md_text)

    # Split lines
    raw_lines = md_text.split('\n')
    i = 0
    total_lines = len(raw_lines)
    
    in_center_div = False
    in_right_div = False
    
    while i < total_lines:
        line = raw_lines[i]
        stripped = line.strip()

        # Handle center div tag
        if '<div align="center">' in stripped.lower():
            in_center_div = True
            i += 1
            continue
        elif '<div align="right">' in stripped.lower():
            in_right_div = True
            i += 1
            continue
        elif '</div>' in stripped.lower():
            in_center_div = False
            in_right_div = False
            i += 1
            continue

        # Handle page breaks
        if '<pagebreak/>' in stripped.lower() or '<pagebreak>' in stripped.lower() or '<pagebreak />' in stripped.lower():
            doc.add_page_break()
            i += 1
            continue

        # Handle HTML table signature block
        if '<table' in stripped.lower():
            # Collect full HTML table
            html_table_lines = []
            while i < total_lines and '</table>' not in raw_lines[i].lower():
                html_table_lines.append(raw_lines[i])
                i += 1
            if i < total_lines:
                html_table_lines.append(raw_lines[i]) # include </table>
                i += 1
            
            table_html = '\n'.join(html_table_lines)
            
            # Check if this is the Guide/HOD or Examiner signature table
            tds = re.findall(r'<td[^>]*>([\s\S]*?)</td>', table_html, flags=re.IGNORECASE)
            if len(tds) == 2:
                left_raw = tds[0].strip()
                right_raw = tds[1].strip()
                
                left_lines = [l.strip() for l in re.split(r'<br\s*/?>|\n', left_raw) if l.strip()]
                right_lines = [l.strip() for l in re.split(r'<br\s*/?>|\n', right_raw) if l.strip()]
                
                add_signature_table(doc, left_lines, right_lines)
            continue

        # Blank line
        if not stripped:
            i += 1
            continue

        # Handle pure <br/> or multiple <br/>
        if re.match(r'^(<br\s*/?>\s*)+$', stripped, flags=re.IGNORECASE):
            # Spacing
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.0
            i += 1
            continue

        # Handle Horizontal Rule ---
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            i += 1
            continue

        # Handle Code Block (``` ... ```)
        if stripped.startswith('```'):
            code_lines = []
            lang = stripped[3:].strip()
            i += 1
            while i < total_lines and not raw_lines[i].strip().startswith('```'):
                code_lines.append(raw_lines[i])
                i += 1
            if i < total_lines:
                i += 1 # skip closing ```

            # Create code box (single cell table with grey background)
            code_text = '\n'.join(code_lines)
            if code_text.strip():
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_shading(cell, "F4F6F8")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                # Table cell border (light grey border)
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                    f'<w:left w:val="single" w:sz="18" w:space="0" w:color="003366"/>' # blue accent bar on left
                    f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                    f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                    f'</w:tcBorders>'
                )
                tcPr.append(tcBorders)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run(sanitize_xml_string(code_text))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1F, 0x24, 0x2F)
                
                # Add spacing after table
                post_p = doc.add_paragraph()
                post_p.paragraph_format.space_before = Pt(0)
                post_p.paragraph_format.space_after = Pt(4)
                post_p.paragraph_format.line_spacing = 1.0
            continue

        # Handle Markdown Tables (| col1 | col2 |)
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < total_lines and raw_lines[i].strip().startswith('|') and raw_lines[i].strip().endswith('|'):
                table_lines.append(raw_lines[i].strip())
                i += 1
            
            # Parse table lines
            rows_data = []
            for tline in table_lines:
                # check if separator row (|---|---|)
                if re.match(r'^\|[\s\-:|]+\|$', tline):
                    continue
                cells = [c.strip() for c in tline.split('|')[1:-1]]
                rows_data.append(cells)
            
            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.style = 'Table Grid'
                
                for row_idx, r_data in enumerate(rows_data):
                    row = tbl.rows[row_idx]
                    is_header = (row_idx == 0)
                    for col_idx in range(num_cols):
                        cell = row.cells[col_idx]
                        cell_val = r_data[col_idx] if col_idx < len(r_data) else ""
                        
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        if is_header:
                            set_cell_shading(cell, "EBF1F5")
                        
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.15
                        if is_header:
                            tokens = parse_inline_tokens(cell_val)
                            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=10.5)
                            for r in p.runs:
                                r.bold = True
                        else:
                            tokens = parse_inline_tokens(cell_val)
                            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=10.0)
                
                # Space after table
                post_p = doc.add_paragraph()
                post_p.paragraph_format.space_before = Pt(0)
                post_p.paragraph_format.space_after = Pt(6)
                post_p.paragraph_format.line_spacing = 1.0
            continue

        # Handle Headings
        # H1: # Heading
        m1 = re.match(r'^#\s+(.+)$', stripped)
        if m1:
            h_text = m1.group(1).strip()
            # If major section / chapter, add page break if not already at start
            is_major = any(h_text.upper().startswith(prefix) for prefix in [
                'CHAPTER', 'DECLARATION', 'CERTIFICATE', 'ABSTRACT', 
                'ACKNOWLEDGEMENT', 'LIST OF', 'REFERENCES', 'APPENDIX'
            ])
            if is_major and len(doc.paragraphs) > 5:
                doc.add_page_break()

            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.5
            if in_center_div or is_major:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            tokens = parse_inline_tokens(h_text)
            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=16)
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # H2: ## Heading
        m2 = re.match(r'^##\s+(.+)$', stripped)
        if m2:
            h_text = m2.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
            if in_center_div:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            tokens = parse_inline_tokens(h_text)
            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=14)
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # H3: ### Heading
        m3 = re.match(r'^###\s+(.+)$', stripped)
        if m3:
            h_text = m3.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.5
            if in_center_div:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            tokens = parse_inline_tokens(h_text)
            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=12)
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # H4: #### Heading
        m4 = re.match(r'^####\s+(.+)$', stripped)
        if m4:
            h_text = m4.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            tokens = parse_inline_tokens(h_text)
            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=12)
            for r in p.runs:
                r.bold = True
                r.italic = True
            i += 1
            continue

        # Bullet List (- item or * item)
        m_bullet = re.match(r'^[-*]\s+(.+)$', stripped)
        if m_bullet:
            item_text = m_bullet.group(1).strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            tokens = parse_inline_tokens(item_text)
            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=12)
            i += 1
            continue

        # Numbered List (1. item)
        m_num = re.match(r'^(\d+)\.\s+(.+)$', stripped)
        if m_num:
            item_text = m_num.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            tokens = parse_inline_tokens(item_text)
            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=12)
            i += 1
            continue

        # Blockquote (> text)
        if stripped.startswith('>'):
            quote_text = stripped[1:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
            tokens = parse_inline_tokens(quote_text)
            add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=11)
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        
        if in_right_div:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif in_center_div:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
        tokens = parse_inline_tokens(stripped)
        add_runs_to_paragraph(p, tokens, default_font="Times New Roman", default_size=12)
        i += 1

    print(f"Saving compiled DOCX to {docx_path}...")
    doc.save(docx_path)
    file_size_mb = os.path.getsize(docx_path) / (1024 * 1024)
    print(f"SUCCESS: Generated {docx_path} ({file_size_mb:.2f} MB)")

if __name__ == '__main__':
    md_file = r"C:\Users\Yash\.gemini\antigravity\brain\9a75b6d9-5102-4e8b-8491-542f65d82384\NTS_Project_Report.md"
    out_file = r"C:\PJT\NTS_Project_Report.docx"
    build_docx_report(md_file, out_file)
